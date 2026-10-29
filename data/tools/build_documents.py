"""Build knowledge documents in many formats (ADR-033, ADR-039).

Written content comes from content.py; generated docs come from the catalog and stores seed.
Every non-Markdown file gets a sidecar <file>.meta.yaml with doc metadata.
Output: data/knowledge/{policies,notices,faq,guides,tables,training,notes,compatibility,generated}/
"""
import csv
import io
import json
import random
from email.message import EmailMessage
from email.utils import format_datetime
from datetime import datetime, timezone
from pathlib import Path

import yaml
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (BaseDocTemplate, Frame, FrameBreak, ListFlowable, ListItem, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)
from reportlab.lib import colors

import content as C

ROOT = Path(__file__).resolve().parents[1]
K = ROOT / "knowledge"
SEED = ROOT / "seed"
rnd = random.Random(20261028)


def meta(path, **m):
    m.setdefault("version", 1)
    m.setdefault("effective_date", "2026-01-01")
    m.setdefault("region", "US")
    m.setdefault("brand_id", None)
    Path(str(path) + ".meta.yaml").write_text(yaml.safe_dump(m, sort_keys=False))


def front_matter(**m):
    m.setdefault("version", 1)
    m.setdefault("effective_date", "2026-01-01")
    m.setdefault("region", "US")
    return "---\n" + yaml.safe_dump(m, sort_keys=False) + "---\n\n"


# ------------------------------------------------------------------ Markdown
def to_markdown(blocks):
    out = []
    for kind, p in blocks:
        if kind == "h1":
            out.append(f"# {p}")
        elif kind == "h2":
            out.append(f"## {p}")
        elif kind == "p":
            out.append(p)
        elif kind == "ul":
            out.append("\n".join(f"- {x}" for x in p))
        elif kind == "ol":
            out.append("\n".join(f"{i}. {x}" for i, x in enumerate(p, 1)))
        elif kind == "table":
            out.append("| " + " | ".join(p[0]) + " |\n|" + "---|" * len(p[0]) + "\n"
                       + "\n".join("| " + " | ".join(r) + " |" for r in p[1:]))
    return "\n\n".join(out) + "\n"


# ------------------------------------------------------------------ DOCX
def to_docx(blocks, path, title, author="BuyWise Policy Team"):
    d = Document()
    d.core_properties.title = title
    d.core_properties.author = author
    d.core_properties.subject = "BuyWise customer policy"
    d.styles["Normal"].font.size = Pt(11)
    for kind, p in blocks:
        if kind == "h1":
            d.add_heading(p, level=1)
        elif kind == "h2":
            d.add_heading(p, level=2)
        elif kind == "p":
            d.add_paragraph(p)
        elif kind in ("ul", "ol"):
            for x in p:
                d.add_paragraph(x, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            t = d.add_table(rows=len(p), cols=len(p[0]))
            t.style = "Table Grid"
            for i, row in enumerate(p):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
                    if i == 0:
                        t.cell(i, j).paragraphs[0].runs[0].bold = True
    d.save(path)


# ------------------------------------------------------------------ digital PDF
def to_pdf(blocks, path, title, two_column_after_h2=None):
    """Digital PDF with repeated header/footer and, optionally, a two-column section (reading-order test)."""
    ss = getSampleStyleSheet()
    W, H = LETTER
    m = 0.9 * inch

    def furniture(canv, doc):
        canv.saveState()
        canv.setFont("Helvetica", 8)
        canv.drawString(m, H - 0.55 * inch, f"BuyWise | {title}")
        canv.drawRightString(W - m, 0.5 * inch, f"Page {doc.page} | BuyWise customer policies | Confidential draft removed")
        canv.restoreState()

    one = Frame(m, m, W - 2 * m, H - 2 * m, id="one")
    gap = 0.3 * inch
    colw = (W - 2 * m - gap) / 2
    left = Frame(m, m, colw, H - 2 * m, id="left")
    right = Frame(m + colw + gap, m, colw, H - 2 * m, id="right")
    doc = BaseDocTemplate(str(path), pagesize=LETTER, title=title, author="BuyWise Policy Team",
                          subject="BuyWise customer policy")
    doc.addPageTemplates([PageTemplate(id="One", frames=[one], onPage=furniture),
                          PageTemplate(id="Two", frames=[left, right], onPage=furniture)])
    from reportlab.platypus import NextPageTemplate, PageBreak
    story = []
    for kind, p in blocks:
        if kind == "h2" and two_column_after_h2 and p == two_column_after_h2:
            story += [NextPageTemplate("Two"), PageBreak()]
        if kind == "h1":
            story.append(Paragraph(p, ss["Title"]))
        elif kind == "h2":
            story.append(Paragraph(p, ss["Heading2"]))
        elif kind == "p":
            story.append(Paragraph(p, ss["BodyText"]))
        elif kind in ("ul", "ol"):
            story.append(ListFlowable([ListItem(Paragraph(x, ss["BodyText"])) for x in p],
                                      bulletType="bullet" if kind == "ul" else "1"))
        elif kind == "table":
            t = Table([[Paragraph(c, ss["BodyText"]) for c in r] for r in p], repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
        story.append(Spacer(1, 6))
    doc.build(story)


# ------------------------------------------------------------------ scanned PDF / PNG
def _font(size):
    for f in ["/System/Library/Fonts/Supplemental/Courier New.ttf", "/System/Library/Fonts/Courier.ttc",
              "/Library/Fonts/Arial.ttf"]:
        if Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default(size=size)


def scanned_image(lines, seed):
    """Render text, then degrade it like a phone scan: skew, blur, noise, uneven lighting."""
    r = random.Random(seed)
    img = Image.new("L", (1700, 2200), 250)
    draw = ImageDraw.Draw(img)
    y = 160
    for i, line in enumerate(lines):
        draw.text((140, y), line, fill=20, font=_font(54 if i == 0 else 40))
        y += 110 if i == 0 else 85
    img = img.rotate(r.uniform(-2.5, 2.5), expand=True, fillcolor=235)
    img = img.filter(ImageFilter.GaussianBlur(1.1))
    px = img.load()
    w, h = img.size
    for _ in range(60000):
        x, yy = r.randrange(w), r.randrange(h)
        px[x, yy] = max(0, min(255, px[x, yy] + r.randint(-90, 60)))
    shade = Image.linear_gradient("L").resize((w, h)).point(lambda v: 255 - v // 6)
    return Image.composite(img, shade, Image.new("L", (w, h), 200)).convert("RGB")


# ------------------------------------------------------------------ HTML with boilerplate
BOILER_HEAD = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{title} | BuyWise Help Center</title>
<meta name="description" content="{title}"><meta name="doc_type" content="{doc_type}">
<script>window.analytics = {{track: function(){{}}}}; analytics.track("pageview");</script>
<style>.cookie{{position:fixed;bottom:0}} nav a{{margin:0 8px}}</style></head><body>
<div class="cookie" id="cookie-banner">We use cookies to improve your experience. <button>Accept all</button> <button>Manage</button></div>
<header><a href="/">BuyWise</a><nav><a href="/shop">Shop</a><a href="/deals">Deals</a><a href="/help">Help Center</a><a href="/account">Account</a></nav>
<form role="search"><input placeholder="Search help articles"></form></header>
<div class="promo">{promo}</div>
<main><article>
"""
BOILER_TAIL = """</article>
<aside class="related"><h3>Related articles</h3><ul><li><a href="/help/shipping">Shipping</a></li><li><a href="/help/pickup">Pickup</a></li><li><a href="/help/payment">Payment</a></li></ul>
<p>Was this article helpful? <button>Yes</button> <button>No</button></p></aside></main>
<footer><p>&copy; 2026 BuyWise. All rights reserved.</p><nav><a href="/privacy">Privacy</a><a href="/terms">Terms</a><a href="/careers">Careers</a></nav>
<p>Sign up for our newsletter: <input type="email" placeholder="you@example.com"></p></footer></body></html>
"""


def faq_html(title, qas, doc_type="faq", promo=""):
    body = [f"<h1>{title}</h1>", "<p class='updated'>Last updated October 1, 2026</p>"]
    for q, a in qas:
        body.append(f"<section class='qa'><h2>{q}</h2><p>{a}</p></section>")
    return BOILER_HEAD.format(title=title, doc_type=doc_type, promo=promo) + "\n".join(body) + BOILER_TAIL


def guide_html(title, steps):
    body = [f"<h1>{title}</h1>", "<ol>"] + [f"<li>{s}</li>" for s in steps] + ["</ol>"]
    return BOILER_HEAD.format(title=title, doc_type="how_to", promo="") + "\n".join(body) + BOILER_TAIL


# ------------------------------------------------------------------ catalog-derived FAQs
def product_faqs(products):
    qas = {"phones": [], "other": []}
    for p in products:
        if p["category"] == "accessory":
            continue
        s, m = p["specs"], p["model"]
        bucket = "phones" if p["category"] == "phone" else "other"
        storages = sorted({v.get("storage_gb") for v in p["variants"] if v.get("storage_gb")})
        if storages:
            qas[bucket].append((f"What storage options does {m} come in?",
                                ", ".join(f"{g // 1024} TB" if g >= 1024 else f"{g} GB" for g in storages) + "."))
        qas[bucket].append((f"What colors is {m} available in?", ", ".join(sorted({v['color'] for v in p['variants']})) + "."))
        if s.get("chip"):
            qas[bucket].append((f"What chip does {m} use?", f"{s['chip']}."))
        if s.get("display"):
            qas[bucket].append((f"What display does {m} have?", f"{s['display']}."))
        tags = p.get("compatibility_tags", [])
        if p["category"] == "phone":
            qas[bucket].append((f"Does {m} support MagSafe?",
                                "Yes, it supports MagSafe and Qi2 wireless charging." if "magsafe" in tags
                                else "MagSafe is not listed for this model in our catalog; charge it with USB-C."))
            if s.get("wired_charging"):
                qas[bucket].append((f"How fast does {m} charge with a cable?", f"{s['wired_charging']}."))
        if s.get("connector"):
            qas[bucket].append((f"What port does {m} have?", f"{s['connector']}."))
        if p["availability"] != "available":
            qas[bucket].append((f"Can I buy {m} now?", f"It is available for {p['availability']}."))
    return qas


# ------------------------------------------------------------------ generated docs (from catalog)
def spec_sheet_md(p):
    rows = [(k.replace("_", " ").capitalize(), v) for k, v in p["specs"].items() if v not in (None, "")]
    storages = sorted({v.get("storage_gb") for v in p["variants"] if v.get("storage_gb")})
    sizes = sorted({v["size"] for v in p["variants"] if v.get("size")})
    md = front_matter(doc_type="spec_sheet", product_family=p["family"], model=p["model"],
                      product_id=p["product_id"], brand_id=p["brand_id"], effective_date=p["specs_verified"],
                      generated_from="data/seed/catalog")
    md += f"# {p['model']} - Technical specifications\n\n"
    md += "| Specification | Value |\n|---|---|\n" + "\n".join(f"| {k} | {v} |" for k, v in rows) + "\n"
    if storages:
        md += f"| Storage options | {', '.join(f'{g // 1024} TB' if g >= 1024 else f'{g} GB' for g in storages)} |\n"
    if sizes:
        md += f"| Case sizes | {', '.join(sizes)} |\n"
    md += f"| Colors | {', '.join(sorted({v['color'] for v in p['variants']}))} |\n"
    md += f"| Release date | {p['release_date']} |\n| Availability | {p['availability']} |\n"
    return md


def overview_md(p, comp_by_dev, names):
    s = p["specs"]
    parts = [f"{p['model']} is a {p['category']} in the {p['family']} family"]
    if s.get("chip"):
        parts.append(f"powered by the {s['chip']} chip")
    text = ", ".join(parts) + "."
    if s.get("display"):
        text += f" It has a {s['display']} display."
    if s.get("battery") or s.get("battery_mah"):
        text += f" Battery: {s.get('battery') or str(s['battery_mah']) + ' mAh'}."
    if p["availability"] != "available":
        text += f" Availability: {p['availability']}."
    acc = sorted({names[a] for a in comp_by_dev.get(p["product_id"], [])})
    md = front_matter(doc_type="product_overview", product_family=p["family"], model=p["model"],
                      product_id=p["product_id"], brand_id=p["brand_id"], effective_date=p["specs_verified"],
                      generated_from="data/seed/catalog")
    md += f"# {p['model']} - Overview\n\n{text}\n\n"
    if acc:
        md += "## Compatible accessories\n\n" + "\n".join(f"- {a}" for a in acc) + "\n"
    return md


def store_md(s):
    md = front_matter(doc_type="store_info", store_id=s["store_id"], generated_from="data/seed/stores.json")
    md += (f"# {s['name']}\n\n- Address: {s['address']}, {s['city']}, {s['state']} {s['zip']}\n"
           f"- Hours: Monday-Saturday 10 am-9 pm, Sunday 11 am-6 pm (see holiday notice for exceptions)\n"
           f"- Time zone: {s['timezone']}\n- Same-day delivery radius: {s['same_day_radius_miles']} miles\n"
           f"- In-store pickup: free; last pickup 30 minutes before closing\n")
    return md


def main():
    for d in ["policies", "notices", "faq", "guides", "tables", "training", "notes", "compatibility",
              "generated/spec_sheets", "generated/overviews", "generated/stores"]:
        (K / d).mkdir(parents=True, exist_ok=True)
    counts = {}

    def count(fmt):
        counts[fmt] = counts.get(fmt, 0) + 1

    # --- DOCX
    docx_docs = [("shipping_policy_v1.docx", C.shipping_policy(1), "Shipping and Delivery Policy v1", 1, "2026-01-01"),
                 ("shipping_policy_v2.docx", C.shipping_policy(2), "Shipping and Delivery Policy v2", 2, "2026-11-15"),
                 ("privacy_policy.docx", C.PRIVACY, "Privacy Policy", 1, "2026-01-01"),
                 ("order_changes_and_limits.docx", C.ORDERS, "Order Changes, Cancellations and Limits", 1, "2026-01-01")]
    for fn, blocks, title, ver, eff in docx_docs:
        path = K / "policies" / fn
        to_docx(blocks, path, title)
        meta(path, doc_type="policy", title=title, version=ver, effective_date=eff,
             supersedes="shipping_policy_v1" if fn == "shipping_policy_v2.docx" else None)
        count("docx")

    # --- digital PDF
    for fn, blocks, title, two in [("warranty_policy.pdf", C.WARRANTY, "Warranty Policy", "How to make a warranty claim"),
                                   ("payment_policy.pdf", C.PAYMENT, "Payment Policy", None),
                                   ("in_store_pickup_policy.pdf", C.PICKUP, "In-Store Pickup Policy", "Who can pick up"),
                                   ("pricing_policy.pdf", C.PRICING, "Pricing Policy", None)]:
        path = K / "policies" / fn
        to_pdf(blocks, path, title, two_column_after_h2=two)
        meta(path, doc_type="policy", title=title)
        count("pdf")

    # --- scanned PDF + PNG (image only, no text layer)
    img = scanned_image(C.HOLIDAY_NOTICE, 1)
    img.save(K / "notices" / "holiday_hours_notice_scan.pdf", "PDF", resolution=200)
    meta(K / "notices" / "holiday_hours_notice_scan.pdf", doc_type="notice", title="Holiday hours 2026",
         effective_date="2026-11-26", expires="2027-01-02")
    scanned_image(C.RATE_CARD, 2).save(K / "notices" / "shipping_rate_card_scan.png")
    meta(K / "notices" / "shipping_rate_card_scan.png", doc_type="notice", title="Shipping rate card (policy v1)",
         superseded_on="2026-11-15")
    counts["scanned_pdf"], counts["png"] = 1, 1

    # --- email with DOCX attachment
    msg = EmailMessage()
    msg["From"] = "BuyWise Policy Team <policy-team@buywise.example>"
    msg["To"] = "all-staff@buywise.example"
    msg["Subject"] = "Policy update: free Standard shipping threshold changes on November 15, 2026"
    msg["Date"] = format_datetime(datetime(2026, 10, 20, 9, 30, tzinfo=timezone.utc))
    msg["Message-ID"] = "<policy-update-2026-10-20@buywise.example>"
    msg.set_content(C.POLICY_UPDATE_EMAIL)
    msg.add_alternative("<html><body>" + "".join(f"<p>{p}</p>" for p in C.POLICY_UPDATE_EMAIL.split("\n\n"))
                        + "</body></html>", subtype="html")
    msg.add_attachment((K / "policies" / "shipping_policy_v2.docx").read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                       filename="shipping_policy_v2.docx")
    (K / "notices" / "policy_update_free_shipping.eml").write_bytes(bytes(msg))
    meta(K / "notices" / "policy_update_free_shipping.eml", doc_type="notice", title="Policy update email",
         effective_date="2026-10-20")
    count("eml")

    # --- Markdown policies
    for fn, blocks, dtype in [("returns_and_out_of_scope.md", C.RETURNS, "policy"),
                              ("accounts_and_guest_checkout.md", C.ACCOUNTS, "policy"),
                              ("support_and_handoff.md", C.SUPPORT, "policy")]:
        (K / "policies" / fn).write_text(front_matter(doc_type=dtype, title=blocks[0][1]) + to_markdown(blocks))
        count("md")

    # --- HTML FAQs (policy + catalog-derived), with boilerplate and two price-leak tests
    products = []
    for f in ["phones", "tablets", "laptops", "watches", "audio", "accessories"]:
        products += json.loads((SEED / "catalog" / f"{f}.json").read_text())
    pf = product_faqs(products)
    faq_pages = {f"faq_{k}": (f"{k.capitalize()} questions", v) for k, v in C.FAQ_POLICY.items()}
    faq_pages["faq_products_phones"] = ("Phone questions", pf["phones"] + [C.PRICE_LEAK_ANSWER])
    faq_pages["faq_products_devices"] = ("Tablet, laptop, watch and audio questions", pf["other"])
    n_faq = 0
    for fn, (title, qas) in faq_pages.items():
        promo = C.PRICE_LEAK_BANNER if fn == "faq_products_phones" else ""
        path = K / "faq" / f"{fn}.html"
        path.write_text(faq_html(title, qas, promo=promo))
        meta(path, doc_type="faq", title=title, effective_date="2026-10-01",
             contains_price_leak_test=fn == "faq_products_phones")
        n_faq += len(qas)
        count("html")

    # --- how-to guides: HTML and Markdown
    for slug, (title, steps) in C.GUIDES.items():
        if slug in C.GUIDES_AS_MARKDOWN:
            (K / "guides" / f"{slug}.md").write_text(
                front_matter(doc_type="how_to", title=title) + f"# {title}\n\n"
                + "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1)) + "\n")
            count("md")
        else:
            path = K / "guides" / f"{slug}.html"
            path.write_text(guide_html(title, steps))
            meta(path, doc_type="how_to", title=title)
            count("html")

    # --- spreadsheets and CSV
    stores = json.loads((SEED / "stores.json").read_text())
    tax = json.loads((SEED / "sales_tax_by_state.json").read_text())
    wb = Workbook()
    ws = wb.active
    ws.title = "Delivery methods"
    ws.merge_cells("A1:D1")
    ws["A1"] = "BuyWise delivery methods - policy v1 (effective 2026-01-01)"
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Method", "Region", "Delivery time (business days)", "Cost"])
    for row in [["Standard", "48 contiguous states", "3-5", "Free over $50, else $5.99"],
                ["Express", "48 contiguous states", "1-2", "$14.99"],
                ["Same-day", "Within 20 miles of a store", "0 (order by 12 pm)", "$9.99"],
                ["Standard", "Alaska and Hawaii", "5-8", "Free over $50, else $5.99"],
                ["Express", "Alaska and Hawaii", "2-4", "$14.99"],
                ["Same-day", "Alaska and Hawaii", "Not available", "-"]]:
        ws.append(row)
    ws2 = wb.create_sheet("Sales tax by state")
    ws2.append(["State", "Combined rate (deliveries)"])
    for t in tax:
        ws2.append([t["state"], t["rate"]])
    ws3 = wb.create_sheet("Store tax (pickup)")
    ws3.append(["Store", "State", "Rate"])
    for s in stores:
        ws3.append([s["name"], s["state"], s["sales_tax_rate"]])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
    wb.properties.title = "Shipping rates and sales tax"
    wb.save(K / "tables" / "shipping_rates_and_tax.xlsx")
    meta(K / "tables" / "shipping_rates_and_tax.xlsx", doc_type="table", title="Shipping rates and sales tax")

    wb = Workbook()
    ws = wb.active
    ws.title = "Regular hours"
    ws.append(["Store", "City", "State", "Mon-Sat", "Sun", "Time zone"])
    for s in stores:
        ws.append([s["name"], s["city"], s["state"], "10:00-21:00", "11:00-18:00", s["timezone"]])
    ws2 = wb.create_sheet("Holidays 2026")
    ws2.merge_cells("A1:C1")
    ws2["A1"] = "Holiday hours - all stores"
    ws2.append(["Date", "Holiday", "Hours"])
    for row in [["2026-11-26", "Thanksgiving Day", "Closed"], ["2026-11-27", "Black Friday", "08:00-22:00"],
                ["2026-12-24", "Christmas Eve", "10:00-16:00"], ["2026-12-25", "Christmas Day", "Closed"],
                ["2026-12-31", "New Year's Eve", "10:00-18:00"], ["2027-01-01", "New Year's Day", "Closed"]]:
        ws2.append(row)
    ws2["A1"].alignment = Alignment(horizontal="center")
    wb.properties.title = "Store hours and holidays"
    wb.save(K / "tables" / "store_hours_and_holidays.xlsx")
    meta(K / "tables" / "store_hours_and_holidays.xlsx", doc_type="table", title="Store hours and holidays",
         expires="2027-01-02")
    counts["xlsx"] = 2

    for fn, rows, title in [
        ("order_limits.csv", [["category", "max_per_order", "scope"], ["phone", "2", "per model"],
                              ["tablet", "5", "per category"], ["laptop", "5", "per category"],
                              ["watch", "5", "per category"], ["audio", "5", "per category"],
                              ["accessory", "5", "per category"]], "Order quantity limits"),
        ("order_status_meanings.csv", [["status", "meaning"]] + [[r[0], r[1]] for r in C.ORDERS[-1][1][1:]],
         "Order status meanings"),
        ("support_hours.csv", [["channel", "hours", "time_zone", "outside_hours"],
                               ["chat with staff", "08:00-22:00", "America/New_York", "ticket; reply next day"],
                               ["stores", "Mon-Sat 10:00-21:00; Sun 11:00-18:00", "store local", "closed"]],
         "Support hours")]:
        with open(K / "tables" / fn, "w", newline="") as f:
            csv.writer(f).writerows(rows)
        meta(K / "tables" / fn, doc_type="table", title=title)
        count("csv")

    # --- PPTX training deck with speaker notes and a table slide
    prs = Presentation()
    for title, bullets, notes in C.TRAINING_SLIDES:
        layout = prs.slide_layouts[0 if bullets is C.TRAINING_SLIDES[0][1] else (5 if bullets is None else 1)]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        if bullets and layout == prs.slide_layouts[1]:
            tf = s.placeholders[1].text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b
        elif bullets:
            s.placeholders[1].text = "\n".join(bullets)
        else:
            rows, cols = len(C.TRAINING_TABLE), 2
            tbl = s.shapes.add_table(rows, cols, Inches(1), Inches(1.8), Inches(8), Inches(3)).table
            for i, r in enumerate(C.TRAINING_TABLE):
                for j, v in enumerate(r):
                    tbl.cell(i, j).text = v
        s.notes_slide.notes_text_frame.text = notes
    prs.core_properties.title = "Staff training: store pickup process"
    prs.save(K / "training" / "staff_training_pickup_process.pptx")
    meta(K / "training" / "staff_training_pickup_process.pptx", doc_type="training",
         title="Staff training: store pickup process", audience="staff", effective_date="2026-10-01")
    count("pptx")

    # --- plain text + written compatibility notes
    (K / "notes" / "internal_quick_tips.txt").write_text(C.QUICK_TIPS)
    meta(K / "notes" / "internal_quick_tips.txt", doc_type="internal_note", title="Staff quick tips",
         audience="staff", effective_date="2026-10-01")
    count("txt")
    (K / "compatibility" / "compatibility_notes.md").write_text(
        front_matter(doc_type="compatibility", title="Compatibility notes")
        + "# Compatibility notes\n\n" + "\n\n".join(f"## {h}\n\n{t}" for h, t in C.COMPAT_NOTES) + "\n")
    count("md")

    # --- generated from catalog and stores
    comp = json.loads((SEED / "catalog" / "compatibility.json").read_text())
    names = {p["product_id"]: p["model"] for p in products}
    by_dev = {}
    for r in comp:
        by_dev.setdefault(r["device_id"], []).append(r["accessory_id"])
    gen = 0
    for p in products:
        if p["category"] == "accessory":
            continue
        (K / "generated" / "spec_sheets" / f"{p['product_id']}.md").write_text(spec_sheet_md(p))
        (K / "generated" / "overviews" / f"{p['product_id']}.md").write_text(overview_md(p, by_dev, names))
        gen += 2
    for s in stores:
        (K / "generated" / "stores" / f"{s['store_id']}.md").write_text(store_md(s))
        gen += 1
    counts["generated_md"] = gen
    print("documents by format:", counts)
    print("faq Q&A pairs:", n_faq)


if __name__ == "__main__":
    main()
